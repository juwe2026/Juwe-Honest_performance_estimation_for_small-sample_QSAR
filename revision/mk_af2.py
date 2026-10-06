# -*- coding: utf-8 -*-
"""Zusatzdatei 2 (neu: Additional file 3) fuer die Revision neu erzeugen (G2-20, Pruefung 03.10.2026).

Ausgangspunkt ist die eingereichte Datei (QSAR_selection_leakage/data/RouteB_descriptor_data.xlsx) mit allen Blaettern.
Geaendert wird nur, was die Korrektur von NumStereoCenters (1,8-Cineol 2 -> 0) und die Neuberechnung der Matrizen beruehrt:
  Descriptors        alle 25 Werte mit RDKit 2026.03.6 neu berechnet und gegen die eingereichten Werte geprueft;
                     NumStereoCenters CIP-basiert (useLegacyImplementation=True); eingereichte Werte als Zusatzspalte
  Computation_rules  Zeile NumStereoCenters
  Pearson_Corr       neu aus den gespeicherten Werten
  Spearman_Corr      neu aus den gespeicherten Werten, Gleichstaende mit mittleren Raengen (Gleitkommarauschen entfernt)
  Reduction          unveraendert; geprueft, dass dieselben Paare > 0.7 und dieselbe Reduktion entstehen
  PLS-DA_reduced     12er-Satz (ohne NumStereoCenters) unveraendert, p-Werte mit Gleichstandsbehandlung wie Table 1
  VIP_two_stage, VIP_gt1_list   mit Cineol = 0 neu (out/twostage25_cineole0.json, Algorithmus code/lit_plsda.py)
  Comparison_3_routes, Conclusions   aus den aktuellen Ergebnissen neu erzeugt (EW 03.10.2026), jede Zahl gegen Table 1/2 geprueft
  Info               Revisionsabschnitt
"""
import os, sys, json, datetime, zoneinfo
os.environ.setdefault('QSAR_ROUTEA_XLSX', '/home/claude/rev/QSAR_selection_leakage/data/20260706_31MT_5Bac_AllDes_cleaned_CS1-4.xlsx')
sys.path.insert(0, '/home/claude/rev/QSAR_selection_leakage/code')
import numpy as np, pandas as pd, openpyxl, rdkit
from openpyxl.styles import Font
from rdkit import Chem
import desc_compute as dc
exec(open('/home/claude/rev/pval_tie.py').read().replace("_NULLDIR = '/home/claude/rev/out/nulls'", "_NULLDIR = '/tmp/claude-0/nulls_af2'"))

SRC = '/home/claude/rev/QSAR_selection_leakage/data/RouteB_descriptor_data.xlsx'
RP = '/home/claude/rev/QSAR_selection_leakage/results/raw_permutations/'
TS = datetime.datetime.now(zoneinfo.ZoneInfo('Europe/Berlin')).strftime('%d.%m.%Y %H%M')
OUT = '/mnt/user-data/outputs/Additional_file_2_RouteB_descriptor_data_neu %s.xlsx' % TS
wb = openpyxl.load_workbook(SRC)
FONT = 'Calibri'

# ------------------------------------------------------------------ Descriptors
ws = wb['Descriptors']
hdr = [c.value for c in ws[1]]
names = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
smiles = [ws.cell(r, 2).value for r in range(2, 2 + len(names))]
assert len(names) == 31
DESC = hdr[hdr.index('MolWt'):]
assert len(DESC) == 25
def compute(smi):
    d = dc.compute(smi); m = Chem.MolFromSmiles(smi)
    d['NumStereoCenters'] = len(Chem.FindMolChiralCenters(m, includeUnassigned=True, useLegacyImplementation=True))
    return {k: (round(v, 6) if isinstance(v, float) else v) for k, v in d.items()}
NEW = [compute(s) for s in smiles]
OLD = [[ws.cell(2 + i, hdr.index(d) + 1).value for d in DESC] for i in range(31)]
changed = []
for i in range(31):
    for k, d in enumerate(DESC):
        o, n = float(OLD[i][k]), float(NEW[i][d])
        if abs(o - n) > 1e-6: changed.append((names[i], d, OLD[i][k], NEW[i][d]))
print('RDKit', rdkit.__version__, '| abweichende Werte:', changed)
assert changed == [('1,8-Cineole', 'NumStereoCenters', 2, 0)], changed
for i in range(31):
    for k, d in enumerate(DESC):
        ws.cell(2 + i, hdr.index(d) + 1, value=NEW[i][d])
cnew = len(hdr) + 1
ws.cell(1, cnew, value='NumStereoCenters_submitted').font = Font(name=FONT, bold=True)
for i in range(31):
    ws.cell(2 + i, cnew, value=OLD[i][DESC.index('NumStereoCenters')])
ws.column_dimensions[openpyxl.utils.get_column_letter(cnew)].width = 26
M = pd.DataFrame([[NEW[i][d] for d in DESC] for i in range(31)], columns=DESC, dtype=float)

# ------------------------------------------------------------ Computation_rules
ws = wb['Computation_rules']
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value == 'NumStereoCenters':
        ws.cell(r, 3, value='FindMolChiralCenters(includeUnassigned=True, useLegacyImplementation=True)')
        ws.cell(r, 5, value='Assigned and unassigned CIP stereocentres; ring-bridgehead atoms that are not stereogenic are not counted '
                            '(1,8-cineole = 0). The submitted version used useLegacyImplementation=False (1,8-cineole = 2); '
                            'those values are kept in column NumStereoCenters_submitted of the Descriptors sheet.')
        break
else:
    raise SystemExit('Computation_rules: NumStereoCenters fehlt')

# ------------------------------------------------------- Pearson / Spearman
P = M.corr(method='pearson'); S = M.corr(method='spearman')   # pandas: mittlere Raenge bei Gleichstand
def write_matrix(name, C):
    w = wb[name]; h = [w.cell(1, c).value for c in range(2, w.max_column + 1)]
    rows = [w.cell(r, 1).value for r in range(2, w.max_row + 1)]
    assert sorted(h) == sorted(DESC) and sorted(rows) == sorted(DESC), name
    old = pd.DataFrame([[w.cell(2 + i, 2 + j).value for j in range(25)] for i in range(25)], index=rows, columns=h, dtype=float)
    for i, a in enumerate(rows):
        for j, b in enumerate(h):
            w.cell(2 + i, 2 + j, value=round(float(C.loc[a, b]), 4))
    return old
Po = write_matrix('Pearson_Corr', P); So = write_matrix('Spearman_Corr', S)
pairs = lambda A: {tuple(sorted((a, b))) for a in DESC for b in DESC if a < b and abs(A.loc[a, b]) > 0.7}
assert pairs(Po) == pairs(P) and pairs(So) == pairs(S), 'Paare > 0.7 veraendert'
nP, nS, nU = len(pairs(P)), len(pairs(S)), len(pairs(P) | pairs(S))
nboth = len(pairs(P) & pairs(S))
assert (nU, nboth, nS - nboth, nP - nboth) == (30, 25, 3, 2), (nU, nboth, nP, nS)
dS = (So - S.loc[So.index, So.columns]).abs(); dP = (Po - P.loc[Po.index, Po.columns]).abs()
mxS = dS.stack().idxmax(); print('max. Aenderung Spearman %.4f %s, Pearson %.4f' % (dS.values.max(), mxS, dP.values.max()))

# Verweis "Table S4" war in der Notiz auf zwei Zeilen getrennt ("... matching Table" / "S4). ..."); fuer die Umnummerierung zusammenfuehren
_w = wb['Reduction']; _hit = 0
for r in range(1, _w.max_row):
    a, b = _w.cell(r, 1).value, _w.cell(r + 1, 1).value
    if isinstance(a, str) and isinstance(b, str) and a.endswith(' matching Table') and b.startswith('S4). '):
        _w.cell(r, 1, value=a[:-len(' Table')]); _w.cell(r + 1, 1, value='Table ' + b); _hit += 1
assert _hit == 1, _hit
# Reduktion wie code/reduction.py (combined, Ausnahme seltener Flags zuerst) auf den neuen Matrizen
w = wb['Reduction']
PUB = {}
for r in range(2, w.max_row + 1):
    a, t, p = w.cell(r, 1).value, w.cell(r, 2).value, w.cell(r, 3).value
    if a in DESC: PUB[a] = (t, [x.strip() for x in str(p).split(',')] if p else [])
RARE = [n for n, (t, _) in PUB.items() if 'rare flag' in str(t)]
link = lambda a, b: max(abs(P.loc[a, b]), abs(S.loc[a, b])) > 0.7
left = [n for n in DESC if n not in RARE]; reps = {}
while left:
    deg = {n: sum(1 for m in left if m != n and link(n, m)) for n in left}
    best = max(left, key=lambda n: (deg[n], -DESC.index(n)))
    part = [m for m in left if m != best and link(best, m)]
    reps[best] = part; left = [m for m in left if m != best and m not in part]
for n in RARE: reps[n] = []
assert set(reps) == set(PUB) and all(set(reps[k]) == set(PUB[k][1]) for k in reps), 'Reduktion veraendert'
print('Reduktion unveraendert:', len(reps), 'Repraesentanten (13, davon C_Count/EsterLacton_flag exakt redundant -> 12)')

# ------------------------------------------------------------ PLS-DA_reduced
w = wb['PLS-DA_reduced']; SH = {'H. influenzae': 'HI', 'S. aureus': 'SA', 'S. pneumoniae': 'SPneu', 'S. pyogenes': 'SPyo', 'P. aeruginosa': 'PA'}
plog = []
for r in range(2, w.max_row + 1):
    s = SH.get(w.cell(r, 1).value)
    if not s: continue
    z = np.load(RP + 'permB_12_%s.npz' % s)
    assert abs(w.cell(r, 3).value - float(z['obs_q2'])) < 5e-4 and abs(w.cell(r, 4).value - float(z['obs_auc'])) < 5e-4
    for col, nul, ob in ((5, z['q2s'], float(z['obs_q2'])), (6, z['aucs'], float(z['obs_auc']))):
        p = round(pval(nul, ob), 4)
        if abs(p - w.cell(r, col).value) > 1e-9: plog.append((s, col, w.cell(r, col).value, p))
        w.cell(r, col, value=p)
print('PLS-DA_reduced, p mit Gleichstandsbehandlung geaendert:', plog)
n = w.max_row + 2
w.cell(n, 1, value='p-values count permuted values within 1e-9 of the observed value as ties, as in Table 1 of the manuscript.').font = Font(name=FONT, size=10, italic=True)

# ------------------------------------------------------ VIP_two_stage / VIP_gt1_list
T = json.load(open('/home/claude/rev/out/twostage25_cineole0.json'))
w = wb['VIP_two_stage']
for r in range(2, w.max_row + 1):
    s = w.cell(r, 1).value; t = T[s]
    assert w.cell(r, 3).value == t['nlv']
    old = tuple(w.cell(r, c).value for c in range(4, 9))
    o2 = t['cineole2']; assert old == tuple(round(o2[k], 3) for k in ('q2_naive', 'auc_naive', 'q2_nested', 'auc_nested')) + (o2['n_vip_gt1'],), (s, old)
    o0 = t['cineole0']
    for c, v in zip(range(4, 9), [round(o0[k], 3) for k in ('q2_naive', 'auc_naive', 'q2_nested', 'auc_nested')] + [o0['n_vip_gt1']]):
        w.cell(r, c, value=v)
n = w.max_row + 2
for line in ['Revision: recomputed with the corrected stereocentre count of 1,8-cineole (0 instead of 2); algorithm and nLV unchanged.',
             'Submitted values (Q2 naive / nested): ' + '; '.join('%s %.3f / %.3f' % (s, T[s]['cineole2']['q2_naive'], T[s]['cineole2']['q2_nested']) for s in T) + '.',
             'The naive Q2 remains higher than the nested Q2 for every strain.']:
    w.cell(n, 1, value=line).font = Font(name=FONT, size=10, italic=True); n += 1
assert all(T[s]['cineole0']['q2_naive'] > T[s]['cineole0']['q2_nested'] for s in T)
w = wb['VIP_gt1_list']
for r in range(2, w.max_row + 1):
    s = w.cell(r, 1).value; t = T[s]
    assert w.cell(r, 2).value == t['nlv'] and w.cell(r, 4).value == ', '.join(t['cineole2']['vip_gt1']), s
    w.cell(r, 3, value=t['cineole0']['n_vip_gt1']); w.cell(r, 4, value=', '.join(t['cineole0']['vip_gt1']))
n = w.max_row + 2
w.cell(n, 1, value='Revision: recomputed with the corrected stereocentre count of 1,8-cineole (0 instead of 2).').font = Font(name=FONT, size=10, italic=True)

# ------------------------------------------------------ Comparison_3_routes und Conclusions neu aus den aktuellen Ergebnissen (EW 03.10.2026)
# Die eingereichten Fassungen enthielten Zahlen eines frueheren Stands (z. B. apparentes Q2 0.43 statt 0.46). Jede Zahl wird hier aus
# derselben Quelle gelesen wie Table 1 des Manuskripts; jede Aussage der Schlussfolgerungen wird per assert geprueft.
import pickle, core
REG = []
O_ = '/home/claude/rev/out/'
T1P = json.load(open(O_ + 'table1_pub_exact.json')); T1C = json.load(open(O_ + 'I_fallback_seeds.json'))
OPR = pickle.load(open(RP + 'ordinary_perm_results.pkl', 'rb'))
B25 = json.load(open(O_ + 'B25_cineole0_seed701.json')); T25 = json.load(open(O_ + 'twostage25_cineole0.json'))
SK = [('H. influenzae', 'HI', 'H.I.'), ('S. aureus', 'SA', 'S.A.'), ('S. pneumoniae', 'SPneu', 'S.Pneu'), ('S. pyogenes', 'SPyo', 'S.Pyo'),
      ('P. aeruginosa', 'PA', 'P.A.')]
def tiep(z, m):
    nul, ob = (z['q2s'], float(z['obs_q2'])) if m == 'q2' else (z['aucs'], float(z['obs_auc'])); return pval(nul, ob), ob
def cls(p): return 'significant after Bonferroni correction (alpha = 0.0125)' if p < 0.0125 else ('nominally significant (p < 0.05)' if p < 0.05 else 'not significant')
ROWS = []
for nm, k, lg in SK:
    a2 = T1C.get(k, {}).get('A2') or T1P['A2'][k]
    zb = np.load(RP + 'permB_12_%s.npz' % k); pb, qb = tiep(zb, 'q2')
    app = float(OPR[lg]['A2'][0])
    interp = ('Artefact (3 actives): descriptive only, no p-value interpreted' if k == 'PA' else
              'Route B: %s; Route A2 nested: %s' % (cls(pb), cls(a2['p_q2'])))
    ROWS.append([nm, round(app, 2), round(a2['q2'], 2), round(a2['p_q2'], 4), round(qb, 2), round(pb, 4), interp])
# Gegenpruefung mit Table 1 des Manuskripts (gedruckte Werte)
_t1 = {}
for _r in json.load(open(O_ + 'numbercheck_rev.json'))['rows']:
    if _r['Table'] == 'Table 1': _t1.setdefault(_r['Endpoint'] + '|' + _r['Row'], _r)   # erste Fundstelle = Teil (a), Q2
for (nm, k, lg), r in zip(SK, ROWS):
    if k == 'PA': continue
    c = _t1[k + '|Route A2']; assert c['Printed'].replace('−', '-').startswith('%+.2f (%.4f)' % (r[2], r[3])), (k, c['Printed'])
    c = _t1[k + '|Route B, 12 descriptors, 1 LV']; assert c['Printed'].startswith('%.2f (%.4f)' % (r[4], r[5])), (k, c['Printed'])
w = wb['Comparison_3_routes']
w.delete_rows(1, w.max_row)
HDR = ['Strain', '(1) Route A2, apparent Q2 (selection on all compounds)', '(2) Route A2, nested Q2 (selection inside every fold)',
       '(3) Route A2, nested p(Q2)', '(4) Route B, 12 descriptors, Q2 (nLV = 1, honest)', '(5) Route B, p(Q2)', 'Reading']
for c, h in enumerate(HDR, 1): w.cell(1, c, value=h).font = Font(name=FONT, bold=True)
for i, r in enumerate(ROWS, 2):
    for c, v in enumerate(r, 1): w.cell(i, c, value=v)
n = len(ROWS) + 3
for line in ['Regenerated for the revision from the result files behind Table 1 of the revised manuscript (Route A2 and Route B rows; apparent',
             'values from the archived ordinary permutation runs). The sheet compares three conditions: the 2,208-descriptor pool of Route A2',
             'with the selection made once on all compounds (column 1, optimistic), the same route with the selection repeated inside every',
             'fold (columns 2 and 3, honest), and the 12 literature-derived descriptors of Route B (columns 4 and 5, honest by construction).',
             'p-values: 2,000 permutations, ties counted as in Table 1. The submitted version of this sheet contained values of an earlier stage.']:
    w.cell(n, 1, value=line).font = Font(name=FONT, size=10, italic=True); n += 1
REG.append('Comparison_3_routes')

# Pearson-only-Reduktion (fuer die Robustheitsaussage) und ihr Modellsatz
def _reduce(mode):
    link = lambda a, b: {'combined': max(abs(P.loc[a, b]), abs(S.loc[a, b])), 'pearson': abs(P.loc[a, b])}[mode] > 0.7
    left = [x for x in DESC if x not in RARE]; reps = {}
    while left:
        deg = {x: sum(1 for m in left if m != x and link(x, m)) for x in left}
        best = max(left, key=lambda x: (deg[x], -DESC.index(x)))
        part = [m for m in left if m != best and link(best, m)]
        reps[best] = part; left = [m for m in left if m != best and m not in part]
    for x in RARE: reps[x] = []
    return reps
RC, RPe = _reduce('combined'), _reduce('pearson')
assert len(RC) == 13 and len(RPe) == 14 and set(RPe) - set(RC) == {'O_Count', 'NumHDonors'} and set(RC) - set(RPe) == {'TPSA'}
assert abs(P.loc['NumHDonors', 'TPSA'] - 0.569) < 5e-4 and abs(S.loc['NumHDonors', 'TPSA'] - 0.7227) < 5e-4
SET12 = list(core.RED12); assert len(SET12) == 12 and 'C_Count' not in SET12 and set(SET12) == set(RC) - {'C_Count'}
SETP = [x for x in RPe if x != 'C_Count']; assert len(SETP) == 13
qP = {k: core.q2(core.y_of(lg), core.loo(core.X_of(SETP), core.y_of(lg), 1)) for _, k, lg in SK}
q12 = {k: float(np.load(RP + 'permB_12_%s.npz' % k)['obs_q2']) for _, k, _ in SK}
dmax = max(abs(qP[k] - q12[k]) for k in ('HI', 'SA', 'SPneu', 'SPyo'))
# Zahlen fuer die Schlussfolgerungen
R12 = {k: dict(zip(('q2', 'pq'), (q12[k], tiep(np.load(RP + 'permB_12_%s.npz' % k), 'q2')[0])), pa=tiep(np.load(RP + 'permB_12_%s.npz' % k), 'auc')[0],
              auc=float(np.load(RP + 'permB_12_%s.npz' % k)['obs_auc'])) for _, k, _ in SK}
A2N = {k: (T1C.get(k, {}).get('A2') or T1P['A2'][k]) for _, k, _ in SK}
assert all(R12[k]['pq'] < 0.0125 for k in ('HI', 'SA', 'SPneu')) and 0.0125 < R12['SPyo']['pq'] < 0.05 and R12['SPyo']['pa'] > 0.05
assert R12['SA']['pa'] > 0.0125 and all(R12[k]['pa'] < 0.0125 for k in ('HI', 'SPneu'))
assert A2N['HI']['p_q2'] > 0.05 and A2N['SPyo']['p_q2'] > 0.05 and A2N['SA']['p_q2'] < 0.05 and A2N['SPneu']['p_q2'] < 0.05
assert all(T25[s]['cineole0']['q2_naive'] > T25[s]['cineole0']['q2_nested'] for s in T25)
f2 = lambda v: ('%.2f' % v).replace('-', '−'); f4 = lambda v: '%.4f' % v
CONC = [
 'CONCLUSIONS - literature-based descriptor selection (leakage-free); regenerated for the revision from the current results',
 '',
 '1) Key finding: with the 12 literature-derived descriptors (correlation-based reduction, one exactly redundant pair resolved),',
 '   H. influenzae, S. aureus and S. pneumoniae show an honest, moderate signal already with one latent variable:',
 '   Q2 = %s, %s and %s, p = %s, %s and %s, all below the Bonferroni threshold of 0.0125 (Table 1).' % (
     f2(R12['HI']['q2']), f2(R12['SA']['q2']), f2(R12['SPneu']['q2']), f4(R12['HI']['pq']), f4(R12['SA']['pq']), f4(R12['SPneu']['pq'])),
 '   Under Route A2 (2,208 descriptors, selection inside every fold) the nested Q2 is %s for H. influenzae (p = %s, not significant),' % (
     f2(A2N['HI']['q2']), f4(A2N['HI']['p_q2'])),
 '   %s for S. aureus (p = %s) and %s for S. pneumoniae (p = %s); see sheet Comparison_3_routes.' % (
     f2(A2N['SA']['q2']), f4(A2N['SA']['p_q2']), f2(A2N['SPneu']['q2']), f4(A2N['SPneu']['p_q2'])),
 '2) Diagnostic value: the failure of the 2,208-descriptor pipeline under nesting is not only a missing signal but the optimism of',
 '   supervised selection from many candidates at n = 31; a small, theory-based, leakage-free set makes the moderate signal visible.',
 '3) Why a simple permutation test is honest here: the reduction uses only descriptor-descriptor correlations (no activity labels),',
 '   so there is no supervised selection step to nest.',
 '4) S. pyogenes remains weak with the 12 descriptors (Q2 = %s, p = %s, nominal only; AUC p = %s, not significant), despite the' % (
     f2(R12['SPyo']['q2']), f4(R12['SPyo']['pq']), f4(R12['SPyo']['pa'])),
 '   largest active class (11). With all 25 descriptors it reaches Q2 = %s (p = %s; Table 2): the reduction removed descriptors' % (
     f2(B25['SPyo']['q2']), f4(B25['SPyo']['p_q2'])),
 '   that carry endpoint-specific information for this strain.',
 '5) P. aeruginosa (3 actives) is an overfitting artefact (AUC = %s) and is reported descriptively only.' % ('%.2f' % R12['PA']['auc']),
 '6) Multiplicity: under Bonferroni (alpha = 0.0125, four evaluable endpoints) the Q2 of H. influenzae, S. aureus and S. pneumoniae',
 '   passes; on AUC, S. aureus (p = %s) does not, H. influenzae (p = %s) and S. pneumoniae (p = %s) do.' % (
     f4(R12['SA']['pa']), f4(R12['HI']['pa']), f4(R12['SPneu']['pa'])),
 '',
 'ON THE VIP QUESTION (PLS-DA with many descriptors, then keep only VIP > 1 and refit):',
 '- As a selection method for the performance estimate this is legitimate only if the VIP selection is repeated inside every',
 '  cross-validation fold (nested); otherwise it is the same leakage as with the univariate filter, with VIP as the criterion.',
 '- Sheet VIP_two_stage: the naive Q2 (selection on all data) exceeds the nested Q2 for every strain, e.g. H. influenzae %s -> %s,' % (
     '%.3f' % T25['H.I.']['cineole0']['q2_naive'], '%.3f' % T25['H.I.']['cineole0']['q2_nested']),
 '  S. pneumoniae %s -> %s. That difference is the leakage.' % ('%.3f' % T25['S.Pneu']['cineole0']['q2_naive'], '%.3f' % T25['S.Pneu']['cineole0']['q2_nested']),
 '- Recommendation: use VIP for interpretation, not as a second selection step that flatters the reported performance.',
 '',
 'NOTE on the number of latent variables: nLV = 1 is the conservative baseline fixed a priori. The values in PLS-DA_reduced for the',
 'nLV that maximises Q2 on all compounds contain the optimism of that choice and are an upper bound; the selection-adjusted',
 'estimate is the doubly nested one in the Supplement (number of latent variables).',
 '',
 'ROBUSTNESS OF THE CORRELATION METRIC (Pearson vs. Spearman):',
 '- Of the %d pairs above 0.7, %d are detected by both coefficients, %d by Spearman only and %d by Pearson only.' % (nU, nboth, nS - nboth, nP - nboth),
 '- Decisive Spearman-only case: NumHDonors - TPSA (r = %.2f, rho = %.2f), monotonic but not linear.' % (P.loc['NumHDonors', 'TPSA'], S.loc['NumHDonors', 'TPSA']),
 '- Pearson alone gives %d representatives instead of %d (before resolving the exactly redundant pair; afterwards %d and %d):' % (len(RPe), len(RC), len(SETP), len(SET12)),
 '  the representative of the size/polarity block shifts from TPSA to O_Count and NumHDonors becomes a singleton.',
 '- With the Pearson-only set the one-component LOO Q2 changes by at most %.2f for the four evaluable endpoints, so the result is' % dmax,
 '  robust to the choice of metric.',
]
w = wb['Conclusions']; w.delete_rows(1, w.max_row)
for i, line in enumerate(CONC, 1):
    w.cell(i, 1, value=line).font = Font(name=FONT, size=10, bold=line.isupper() or line.startswith(('CONCLUSIONS', 'ON THE', 'ROBUSTNESS', 'NOTE')))
REG.append('Conclusions')
print('Pearson-only-Satz: Q2 max. Abweichung %.3f' % dmax, {k: round(qP[k], 3) for k in qP})

# ------------------------------------------------------------------ Info
w = wb['Info']
for r in range(1, w.max_row + 1):
    v = w.cell(r, 1).value
    if isinstance(v, str):
        v = (v.replace('See the Descriptors, Computation_rules, Reduction, PLS-DA_reduced, VIP_reduced, VIP_two_stage,',
                       'See the Descriptors, Computation_rules, Pearson_Corr, Spearman_Corr, Reduction, PLS-DA_reduced, VIP_reduced,')
              .replace('Comparison_3_routes, Conclusions and VIP_gt1_list sheets for the data and results themselves.',
                       'VIP_two_stage, Comparison_3_routes, Conclusions and VIP_gt1_list sheets for the data and results themselves.')
              .replace('section of the manuscript and the NumStereoCenters column, which counts unassigned potential',
                       'section of the manuscript and the NumStereoCenters column, which counts assigned and unassigned CIP')
              .replace('stereocentres. For the same reason', 'stereocentres. For the same reason'))
        w.cell(r, 1, value=v)
txt = [w.cell(r, 1).value for r in range(1, w.max_row + 1)]
assert any('Conclusions and VIP_gt1_list sheets' in str(x) for x in txt) and any('assigned and unassigned CIP' in str(x) for x in txt)
n = w.max_row + 2
REV = [
 'REVISION (%s)' % TS.split()[0],
 'Software: RDKit %s, Python %s. The RDKit version used for the submitted file was not recorded.' % (rdkit.__version__, sys.version.split()[0]),
 'All 25 descriptor values were recomputed with this version from the input SMILES. Every submitted value was reproduced',
 'except one: NumStereoCenters of 1,8-cineole, which the submitted file gave as 2. That file had used',
 'FindMolChiralCenters(..., useLegacyImplementation=False), which also counts ring-bridgehead atoms as potential',
 'stereo elements; the CIP-based count (useLegacyImplementation=True, the RDKit default) gives 0, the correct value for',
 'this achiral molecule. The corrected value is used here and in the revised manuscript; the submitted values are kept',
 'in column NumStereoCenters_submitted. Floating-point values are stored to six decimals.',
 'Pearson_Corr and Spearman_Corr were recomputed from the stored values. Spearman coefficients use average ranks for ties.',
 'In the submitted file they had been computed from unrounded values, in which isomers of the same formula differ in MolWt',
 'and LabuteASA in the 14th decimal; this broke ties and caused the deviations of up to 0.019 noted in review.',
 'Largest change of a Spearman coefficient against the submitted sheet: %.4f (%s / %s), including the effect of the' % (dS.values.max(), mxS[0], mxS[1]),
 'corrected stereocentre count. The %d pairs with |r| or |rho| > 0.7 (%d by both coefficients, %d by Spearman only, %d by' % (nU, nboth, nS - nboth, nP - nboth),
 'Pearson only) and the reduction in the Reduction sheet are unchanged.',
 'NumStereoCenters is not among the 12 modelled descriptors, so PLS-DA_reduced and VIP_reduced are unchanged apart from',
 'p-values, which now count ties as in Table 1. VIP_two_stage and VIP_gt1_list, which use all 25 descriptors, were',
 'recomputed with the corrected value.',
 'The summary sheets Comparison_3_routes and Conclusions were regenerated from the current results (the submitted versions',
 'contained values of an earlier stage); every number in them is read from the same result files as Table 1 and Table 2.',
]
for line in REV:
    w.cell(n, 1, value=line).font = Font(name=FONT, size=10, bold=line.startswith('REVISION')); n += 1
wb.save(OUT)
print('geschrieben:', OUT)
json.dump(dict(rdkit=rdkit.__version__, changed=changed, max_dS=float(dS.values.max()), max_dS_pair=list(mxS), max_dP=float(dP.values.max()),
               pairs=dict(union=nU, both=nboth, spearman_only=nS - nboth, pearson_only=nP - nboth), pls_p_changes=plog, regenerated=REG, out=OUT),
          open('/home/claude/rev/out/af2_regen.json', 'w'), indent=1)
