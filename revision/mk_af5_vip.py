# -*- coding: utf-8 -*-
"""Zusatzdatei 5 (neu: Additional file 8, VIP- und Effektstaerke-Tabellen) fuer die Revision (G2-20, 03.10.2026).
Nur die Spalten der Blaetter VIP_Lit_<Stamm>, die alle 25 Literaturdeskriptoren verwenden (TwoStage_VIP_apparent,
TwoStage_VIP_nested_mean, TwoStage_selfreq_VIPgt1, EffectSize), werden mit NumStereoCenters von 1,8-Cineol = 0 neu
berechnet (out/vip25_cineole0.json, Algorithmus code/compute_vip.py, zwei Komponenten). Die RouteB_*-Spalten (12er-Satz,
ohne NumStereoCenters) und die ChemDes-Blaetter bleiben unveraendert. Vorher wird geprueft, dass die eingereichten Werte
mit Cineol = 2 reproduziert werden. Zeilen werden wie in der Einreichung nach TwoStage_VIP_apparent absteigend sortiert."""
import os, sys, json, datetime, zoneinfo, numpy as np, openpyxl, pandas as pd
os.environ.setdefault('QSAR_ROUTEA_XLSX', '/home/claude/rev/QSAR_selection_leakage/data/20260706_31MT_5Bac_AllDes_cleaned_CS1-4.xlsx')
sys.path.insert(0, '/home/claude/rev/QSAR_selection_leakage/code')
import core, desc_compute as dc
from scipy.stats import rankdata
from openpyxl.styles import Font
SRC = '/mnt/user-data/uploads/J_Chemometrics/2026-09-01 Manuskript+Suppl+addfile2-10 upload/2026-09-01 1805 Additional_file_5_VIP_and_effect_size_tables.xlsx'
TS = datetime.datetime.now(zoneinfo.ZoneInfo('Europe/Berlin')).strftime('%d.%m.%Y %H%M')
OUT = '/mnt/user-data/outputs/Additional_file_5_VIP_and_effect_size_tables_neu %s.xlsx' % TS
V = json.load(open('/home/claude/rev/out/vip25_cineole0.json')); L25 = V['L25']
X2 = core.X_of(L25); ci = core.COMPOUNDS.index('1,8-Cineole'); j = L25.index('NumStereoCenters')
assert X2[ci, j] == 2
X0 = X2.copy(); X0[ci, j] = 0
# Die eingereichte Effektstaerke von MolWt wurde aus ungerundeten Werten berechnet, in denen Isomere gleicher Summenformel
# in der 14. Stelle differieren (Gleichstaende aufgebrochen, vgl. Spearman-Matrix); hier mit Gleichstaenden.
_D = pd.read_excel('/home/claude/rev/QSAR_selection_leakage/data/RouteB_descriptor_data.xlsx', 'Descriptors')
_smi = dict(zip(_D.Monoterpenoid, _D.SMILES))
MW_NOISY = np.array([dc.compute(_smi[c])['MolWt'] for c in core.COMPOUNDS])
assert np.allclose(MW_NOISY, X2[:, L25.index('MolWt')], atol=1e-9)
def eff_rb(x, y):
    a = x[y == 1]; b = x[y == 0]
    if len(np.unique(np.concatenate([a, b]))) < 2: return 0.0
    r = rankdata(np.concatenate([a, b])); U1 = r[:len(a)].sum() - len(a) * (len(a) + 1) / 2
    return abs(2 * U1 / (len(a) * len(b)) - 1)
wb = openpyxl.load_workbook(SRC); LOG = []
for s in core.STRAINS:
    ws = wb['VIP_Lit_' + s]; y = core.y_of(s); R = V['res'][s]
    hdr = [c.value for c in ws[1]]
    assert hdr == ['Descriptor', 'RouteB_VIP_apparent', 'RouteB_VIP_nested_mean', 'TwoStage_VIP_apparent', 'TwoStage_VIP_nested_mean',
                   'TwoStage_selfreq_VIPgt1', 'EffectSize'], hdr
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2) if r[0].value]
    assert sorted(r[0] for r in rows) == sorted(L25)
    new = []
    for r in rows:
        k = L25.index(r[0])
        assert abs(r[3] - R['app2'][k]) < 6e-4 and abs(r[4] - R['nest2'][k]) < 6e-4 and r[5] == '%d/31' % R['freq2'][k], (s, r)
        assert abs(r[6] - eff_rb(MW_NOISY if r[0] == 'MolWt' else X2[:, k], y)) < 6e-4, (s, r[0], r[6])
        n = r[:3] + [round(R['app'][k], 3), round(R['nest'][k], 3), '%d/31' % R['freq'][k], round(eff_rb(X0[:, k], y), 3)]
        if n != r: LOG.append((s, r[0], r[3:], n[3:]))
        new.append(n)
    new.sort(key=lambda r: (-r[3], r[0]))
    for i, r in enumerate(new):
        for c, v in enumerate(r): ws.cell(2 + i, 1 + c, value=v)
w = wb['README']; n = w.max_row + 2
for line in ['REVISION (%s): in the VIP_Lit sheets the columns TwoStage_VIP_apparent, TwoStage_VIP_nested_mean, TwoStage_selfreq_VIPgt1 and EffectSize, '
             'which use all 25 literature descriptors, were recomputed with the corrected stereocentre count of 1,8-cineole (0 instead of 2; see the '
             'Route B descriptor file). Rows are again sorted by TwoStage_VIP_apparent. The RouteB columns (12-descriptor set, which does not contain '
             'NumStereoCenters) and the ChemDes sheets are unchanged. The EffectSize of MolWt is now computed with ties between isomers of the same '
             'formula, which the submitted values had broken because of floating-point noise in the 14th decimal (change at most 0.024).' % TS.split()[0]]:
    w.cell(n, 1, value=line).font = Font(bold=True); n += 1
wb.save(OUT)
ns = [l for l in LOG if l[1] in ('NumStereoCenters', 'MolWt')]
assert max(abs(l[2][3] - l[3][3]) for l in LOG if l[1] == 'MolWt') <= 0.024
mx = max(abs(a - b) for l in LOG if l[1] != 'NumStereoCenters' for a, b in zip(l[2][:2], l[3][:2]))
print('geaendert: %d Zeilen; groesste Aenderung ausser NumStereoCenters: %.3f' % (len(LOG), mx))
for l in ns: print('  ', l)
json.dump(dict(out=OUT, changes=LOG, max_other=mx), open('/home/claude/rev/out/af5_vip_regen.json', 'w'), indent=1, default=str)
print('geschrieben:', OUT)
