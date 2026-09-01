"""Add machine-readable structure identifiers to Additional file 2.

Writes three columns into the ``Descriptors`` sheet, immediately after the existing
``SMILES`` column:

    Canonical_SMILES   RDKit canonical SMILES
    InChI              standard InChI
    InChIKey           standard InChIKey

All three are derived from the string already in the ``SMILES`` column, that is from the
exact structure the Route B descriptors were computed from by ``code/desc_compute.py``.
Nothing is looked up externally and no structure is edited, so the identifiers cannot
disagree with the descriptor values beside them; the script verifies this by recomputing
every descriptor from the same SMILES and comparing it with the value in the sheet before
it writes anything.

Two consequences of that choice are worth stating, and are also written into the file's
Info sheet:

* The modelled structures carry no stereochemistry except for the two double-bond
  geometries specified in the input (Citral, Geranic acid). The identifiers therefore
  describe the constitution that was modelled, not the configuration of the material
  that was tested, whose stereochemistry the compound names give. This matches the
  Methods section, which states that descriptors were computed from neutral, canonical
  SMILES without 3D optimisation and that ``NumStereoCenters`` counts unassigned
  potential stereocentres.
* Because they pin the constitution exactly, the identifiers resolve the one compound
  name in the series that is ambiguous: "3-Carene, alpha-Carene" was modelled as
  3-carene (car-3-ene), InChIKey BQOFWKZOCNGFEC-UHFFFAOYSA-N.

Usage:  python code/add_structure_identifiers.py IN.xlsx [OUT.xlsx]

Run against the bundled ``data/RouteB_descriptor_data.xlsx``, which already carries the three
columns, the script re-derives them and checks them against what is in the file instead of
writing, so the published identifiers can be verified from this archive in one command:

    python code/add_structure_identifiers.py data/RouteB_descriptor_data.xlsx
"""
import sys
import openpyxl
from rdkit import Chem, RDLogger

RDLogger.DisableLog("rdApp.*")

NEW_COLS = ["Canonical_SMILES", "InChI", "InChIKey"]

RULES = [
    ("Canonical_SMILES", "Structure identifier", "RDKit Chem.MolToSmiles", None,
     "RDKit canonical SMILES of the structure in the SMILES column"),
    ("InChI", "Structure identifier", "RDKit Chem.MolToInchi", None,
     "Standard InChI (prefix InChI=1S) of the same structure"),
    ("InChIKey", "Structure identifier", "RDKit Chem.InchiToInchiKey", None,
     "Standard InChIKey of the same structure"),
]

INFO = [
    "",
    "Machine-readable structure identifiers. The Descriptors sheet gives, for every compound, the input",
    "SMILES together with its RDKit canonical SMILES, standard InChI and standard InChIKey. All three are",
    "derived from that same input SMILES, that is from the exact structure the descriptors in this file were",
    "computed from, so they identify the modelled structure and not a looked-up reference structure.",
    "The modelled structures are neutral and two-dimensional and carry no stereochemistry, apart from the",
    "double-bond geometry specified for Citral and Geranic acid; the compound names give the configuration of",
    "the material that was tested. The identifiers therefore describe constitution, matching the Methods",
    "section of the manuscript and the NumStereoCenters column, which counts unassigned potential",
    "stereocentres. For the same reason the InChIKey resolves the one ambiguous name in the series:",
    "\"3-Carene, alpha-Carene\" was modelled as 3-carene (car-3-ene), BQOFWKZOCNGFEC-UHFFFAOYSA-N.",
]


def identifiers(smiles):
    mol = Chem.MolFromSmiles(smiles)
    assert mol is not None, "unparsable SMILES: %r" % (smiles,)
    inchi = Chem.MolToInchi(mol)
    return Chem.MolToSmiles(mol), inchi, Chem.InchiToInchiKey(inchi)


def verify(ws):
    """Recompute every Route B descriptor from the SMILES column and compare."""
    import os
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import desc_compute as dc
    rows = list(ws.values)
    hdr = list(rows[0])
    n = 0
    for r in [r for r in rows[1:] if r[0]]:
        for k, v in dc.compute(r[1]).items():
            ref = r[hdr.index(k)]
            assert abs(float(v) - float(ref)) <= 5e-4, (r[0], k, v, ref)
            n += 1
    return n


def check(ws, hdr):
    """The columns are already there: re-derive and compare instead of writing."""
    bad = 0
    for i in range(2, ws.max_row + 1):
        smi = ws.cell(row=i, column=2).value
        if not smi:
            continue
        for name, want in zip(NEW_COLS, identifiers(smi)):
            got = ws.cell(row=i, column=hdr.index(name) + 1).value
            if got != want:
                bad += 1
                print("MISMATCH %s %s\n  in file: %s\n  derived: %s"
                      % (ws.cell(row=i, column=1).value, name, got, want))
    print("identifiers already present: re-derived and compared, %d mismatches" % bad)
    return bad


def main(src, dst=None):
    wb = openpyxl.load_workbook(src)
    ws = wb["Descriptors"]
    hdr = [c.value for c in ws[1]]
    assert hdr[:2] == ["Monoterpenoid", "SMILES"], hdr[:2]

    n = verify(ws)
    print("verified %d descriptor values against the SMILES column" % n)

    if all(c in hdr for c in NEW_COLS):
        raise SystemExit(1 if check(ws, hdr) else 0)
    assert not any(c in hdr for c in NEW_COLS), "identifier columns only partly present"
    assert dst, "give an output path to write the identifier columns"

    ws.insert_cols(3, len(NEW_COLS))
    for j, name in enumerate(NEW_COLS):
        ws.cell(row=1, column=3 + j, value=name)
    keys = {}
    for i in range(2, ws.max_row + 1):
        smi = ws.cell(row=i, column=2).value
        if not smi:
            continue
        vals = identifiers(smi)
        for j, v in enumerate(vals):
            ws.cell(row=i, column=3 + j, value=v)
        keys.setdefault(vals[2], []).append(ws.cell(row=i, column=1).value)
    dup = {k: v for k, v in keys.items() if len(v) > 1}
    assert not dup, dup
    print("wrote identifiers for %d compounds, all InChIKeys distinct" % len(keys))

    rules = wb["Computation_rules"]
    for row in RULES:
        rules.append(list(row))
    info = wb["Info"]
    for line in INFO:
        info.append([line])

    wb.save(dst)
    print("saved %s" % dst)


if __name__ == "__main__":
    main(*sys.argv[1:3])
