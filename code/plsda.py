import numpy as np, pandas as pd
from sklearn.cross_decomposition import PLSRegression
from sklearn.metrics import roc_auc_score, matthews_corrcoef
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def autoscale(Xtr, Xte):
    mu = Xtr.mean(0); sd = Xtr.std(0, ddof=1); sd[sd == 0] = 1.0
    return (Xtr - mu) / sd, (Xte - mu) / sd

def loo_cont(X, y, nlv):
    n = len(y); yhat = np.zeros(n)
    for i in range(n):
        tr = np.ones(n, bool); tr[i] = False
        Xtr, Xte = autoscale(X[tr], X[i:i+1])
        m = PLSRegression(n_components=nlv, scale=False).fit(Xtr, y[tr])
        yhat[i] = m.predict(Xte).ravel()[0]
    return yhat

def q2(y, yh): return 1 - np.sum((y - yh)**2) / np.sum((y - y.mean())**2)

def r2y_apparent(X, y, nlv):
    Xs = (X - X.mean(0)) / np.where(X.std(0, ddof=1) == 0, 1, X.std(0, ddof=1))
    m = PLSRegression(n_components=nlv, scale=False).fit(Xs, y)
    yh = m.predict(Xs).ravel()
    return 1 - np.sum((y - yh)**2) / np.sum((y - y.mean())**2)

def fit_full(X, y, nlv):
    mu = X.mean(0); sd = X.std(0, ddof=1); sd[sd == 0] = 1.0
    Xs = (X - mu) / sd
    m = PLSRegression(n_components=nlv, scale=False).fit(Xs, y)
    yh = m.predict(Xs).ravel()
    r2y = 1 - np.sum((y - yh)**2) / np.sum((y - y.mean())**2)
    # R2X kumulativ
    T = m.x_scores_; P = m.x_loadings_
    ssx = np.sum(Xs**2)
    r2x = np.sum([(T[:, a]**2).sum() * (P[:, a]**2).sum() for a in range(nlv)]) / ssx
    # VIP
    W = m.x_weights_; Q = m.y_loadings_.ravel()
    ssy_a = np.array([(Q[a]**2) * (T[:, a]**2).sum() for a in range(nlv)])
    p = X.shape[1]; vip = np.zeros(p)
    for j in range(p):
        wnorm = (W[j, :] / np.linalg.norm(W, axis=0))**2
        vip[j] = np.sqrt(p * np.sum(ssy_a * wnorm) / ssy_a.sum())
    return r2y, r2x, vip

def classification(y, yh, thr=0.5):
    pred = (yh >= thr).astype(int)
    TP = int(((pred == 1) & (y == 1)).sum()); FN = int(((pred == 0) & (y == 1)).sum())
    FP = int(((pred == 1) & (y == 0)).sum()); TN = int(((pred == 0) & (y == 0)).sum())
    sens = TP / (TP + FN) if (TP + FN) else np.nan
    spec = TN / (TN + FP) if (TN + FP) else np.nan
    prec = TP / (TP + FP) if (TP + FP) else np.nan
    acc = (TP + TN) / len(y); ba = np.nanmean([sens, spec])
    try: mcc = matthews_corrcoef(y, pred)
    except Exception: mcc = np.nan
    return dict(TP=TP, FN=FN, FP=FP, TN=TN, sens=sens, spec=spec, prec=prec, acc=acc, ba=ba, mcc=mcc)

def run_endpoint(fpath, short, full, nperm=2000, seed=42, maxlv=6):
    rng = np.random.default_rng(seed)
    xl = pd.ExcelFile(fpath)
    fin = pd.read_excel(xl, f'FIN_Des_{short}')
    cs = pd.read_excel(xl, f'CS1-5_{short}')
    reps = fin['Representative'].tolist()
    es_vals = fin['ES_Rep'].tolist()
    ycol = cs.columns[2]
    X = cs[reps].apply(pd.to_numeric, errors='coerce').values.astype(float)
    y = (cs[ycol].astype(str).str.strip() == 'active').astype(float).values
    n = len(y); n1 = int(y.sum()); n0 = int(n - n1)

    # Komponentenwahl: max LOO-Q2 (Tie: kleineres nLV)
    scan = []
    for nlv in range(1, min(maxlv, X.shape[1], n - 2) + 1):
        yh = loo_cont(X, y, nlv); scan.append((nlv, q2(y, yh), roc_auc_score(y, yh)))
    nlv = sorted(scan, key=lambda r: (-r[1], r[0]))[0][0]

    yh = loo_cont(X, y, nlv)
    Q2 = q2(y, yh); AUC = roc_auc_score(y, yh); cls = classification(y, yh)
    R2Y, R2X, vip = fit_full(X, y, nlv)

    # Permutationstest
    permQ = np.empty(nperm); permA = np.empty(nperm); permR = np.empty(nperm)
    for k in range(nperm):
        yp = rng.permutation(y)
        yhp = loo_cont(X, yp, nlv)
        permQ[k] = q2(yp, yhp); permA[k] = roc_auc_score(yp, yhp)
        permR[k] = r2y_apparent(X, yp, nlv)
    pQ = (1 + np.sum(permQ >= Q2)) / (1 + nperm)
    pA = (1 + np.sum(permA >= AUC)) / (1 + nperm)
    pR = (1 + np.sum(permR >= R2Y)) / (1 + nperm)

    stats = dict(short=short, full=full, n=n, n1=n1, n0=n0, nlv=nlv, scan=scan,
                 R2X=R2X, R2Y=R2Y, Q2=Q2, AUC=AUC, cls=cls,
                 permQ=(permQ.mean(), pQ), permA=(permA.mean(), pA), permR=(permR.mean(), pR),
                 reps=reps, vip=vip, es=es_vals, nperm=nperm)
    write_sheet(fpath, stats)
    return stats

def write_sheet(fpath, s):
    wb = openpyxl.load_workbook(fpath)
    name = f'PLS-DA_{s["short"]}'
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name)
    H = Font(bold=True, size=12); B = Font(bold=True)
    fill = PatternFill('solid', fgColor='DDE7F0')
    r = 1
    def put(row, vals, bold=False, section=False):
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=j, value=v)
            if bold: c.font = B
            if section: c.font = H; c.fill = fill
        return row + 1

    r = put(r, [f'PLS-DA – {s["full"]} ({s["short"]})'], section=True)
    r = put(r, [f'Erstellt aus den finalen Deskriptoren (Repraesentanten der Feature Groups + Singletons) aus FIN_Des_{s["short"]}'])
    r += 1
    r = put(r, ['MODELLPARAMETER'], section=True)
    r = put(r, ['Latente Variablen (nLV)', s['nlv']])
    r = put(r, ['Deskriptoren (Input)', len(s['reps'])])
    r = put(r, ['Objekte', f'{s["n"]} (active={s["n1"]}, non-active={s["n0"]})'])
    r = put(r, ['Skalierung', 'Autoscaling (Z-Score), pro CV-Fold neu geschaetzt'])
    r = put(r, ['CV-Methode', 'Leave-One-Out (LOO)'])
    r = put(r, ['Komponentenwahl', 'max. LOO-Q2 (Tie-break: weniger Komponenten)'])
    r += 1
    r = put(r, ['GUETE / VORHERSAGE (LOO-kreuzvalidiert, sofern nicht anders vermerkt)'], section=True)
    r = put(r, ['R2X (kumulativ)', round(s['R2X'], 3)])
    r = put(r, ['R2Y (apparent, Full-Fit)', round(s['R2Y'], 3)])
    r = put(r, ['Q2 (LOO)', round(s['Q2'], 3)])
    r = put(r, ['AUC (LOO)', round(s['AUC'], 3)])
    c = s['cls']
    r = put(r, ['Balanced Accuracy (thr=0.5)', round(c['ba'], 3)])
    r = put(r, ['Sensitivitaet / active (thr=0.5)', round(c['sens'], 3)])
    r = put(r, ['Spezifitaet (thr=0.5)', round(c['spec'], 3)])
    r = put(r, ['Precision / active (thr=0.5)', None if np.isnan(c['prec']) else round(c['prec'], 3)])
    r = put(r, ['Accuracy (thr=0.5)', round(c['acc'], 3)])
    r = put(r, ['MCC (thr=0.5)', round(c['mcc'], 3)])
    r = put(r, ['Konfusion thr=0.5 (TP/FN/FP/TN)', f'{c["TP"]}/{c["FN"]}/{c["FP"]}/{c["TN"]}'])
    r += 1
    r = put(r, [f'PERMUTATIONSTEST ({s["nperm"]} Permutationen der Y-Labels, nLV fix)'], section=True)
    r = put(r, ['Statistik', 'Beobachtet', 'Perm-Mittel', 'p-Wert'], bold=True)
    r = put(r, ['Q2 (LOO)', round(s['Q2'], 3), round(s['permQ'][0], 3), round(s['permQ'][1], 4)])
    r = put(r, ['AUC (LOO)', round(s['AUC'], 3), round(s['permA'][0], 3), round(s['permA'][1], 4)])
    r = put(r, ['R2Y (apparent)', round(s['R2Y'], 3), round(s['permR'][0], 3), round(s['permR'][1], 4)])
    r = put(r, ['p = (1 + #{Perm >= Beobachtet}) / (1 + n_perm)'])
    r += 1
    r = put(r, ['VIP (Variable Importance in Projection), absteigend'], section=True)
    r = put(r, ['Descriptor', 'VIP', 'EffectSize (Step5)'], bold=True)
    order = np.argsort(-s['vip'])
    for idx in order:
        r = put(r, [s['reps'][idx], round(float(s['vip'][idx]), 3),
                    None if s['es'][idx] != s['es'][idx] else round(float(s['es'][idx]), 3)])
    r += 1
    r = put(r, ['KOMMENTARE'], section=True)
    for line in build_comments(s):
        cc = ws.cell(row=r, column=1, value=line); cc.alignment = Alignment(wrap_text=False); r += 1

    ws.column_dimensions['A'].width = 40
    for col in ['B', 'C', 'D']: ws.column_dimensions[col].width = 16
    wb.save(fpath)

def build_comments(s):
    Q2, AUC, p_q = s['Q2'], s['AUC'], s['permQ'][1]
    L = []
    L.append(f'Interpretation ({s["full"]}):')
    if s['n1'] <= 3:
        L.append(f'!!! WARNUNG: Nur {s["n1"]} aktive Objekte. Der Wilcoxon-Filter (Step 5) und diese PLS-DA')
        L.append('    sind hier statistisch NICHT belastbar. Alle Kennzahlen nur als Formalitaet zu lesen,')
        L.append('    nicht als Evidenz fuer ein funktionierendes Modell.')
    if Q2 > 0.4 and p_q < 0.05:
        L.append(f'- Q2={Q2:.2f} (>0) und Permutations-p={p_q:.3f} deuten auf ueberzufaellige Trennung hin.')
    elif Q2 > 0:
        L.append(f'- Q2={Q2:.2f} ist positiv, aber schwach; Permutations-p(Q2)={p_q:.3f}. Vorsicht bei der Deutung.')
    else:
        L.append(f'- Q2={Q2:.2f} ist <=0: das Modell sagt kreuzvalidiert schlechter voraus als der Mittelwert.')
        L.append('  Trotz evtl. hoher apparenter R2Y/AUC also KEINE belastbare Vorhersagekraft.')
    L.append(f'- AUC(LOO)={AUC:.2f}; wegen der Klassenimbalance sind AUC/Balanced Accuracy/MCC aussagekraeftiger als Accuracy.')
    L.append('')
    L.append('Wichtige methodische Vorbehalte:')
    L.append('- SELECTION BIAS: Die Deskriptoren wurden in Step 5 supervidiert (Effektgroesse+p) auf dem GESAMT-')
    L.append('  datensatz ausgewaehlt. Diese CV-/Permutationswerte sind daher OPTIMISTISCH verzerrt: der Permutations-')
    L.append('  test variiert nur die Modellbildung, nicht die vorgelagerte Merkmalsauswahl.')
    L.append('- Ehrlicher waere ein VOLLSTAENDIG genesteter Ablauf: In jeder CV-Fold / jeder Permutation die Effekt-')
    L.append('  groessen-Selektion + Feature-Group-Bildung neu auf nur den Trainingsdaten durchfuehren.')
    L.append(f'- n={s["n"]} bei {len(s["reps"])} Deskriptoren (p ~ n). Ergebnisse sind instabil; externe Validierung fehlt.')
    L.append('- Empfehlung Step 6: nested CV (Selektion IN den Folds), zusaetzlich Bootstrap/Y-Scrambling der Gesamtpipeline.')
    return L
