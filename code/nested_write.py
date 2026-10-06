# Revision 03.10.2026: Permutations-p-Werte zaehlen Gleichstaende mit (Toleranz 1e-9), siehe CHANGELOG_revision.md
import numpy as np, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill
from collections import Counter
import nested_generic as ng

BIASED={'P.A.':(0.88,1.00),'S.A.':(0.50,0.97),'S. Pneu':(0.56,0.97),'S. Pyo':(0.52,0.95)}
EP=[('P.A.','P. aeruginosa','CS1-4_P.A.','PA','/mnt/user-data/outputs/20260706_31MT_P.A._Step5.xlsx'),
    ('S.A.','Staph. aureus','CS1-4_S.A.','SA','/mnt/user-data/outputs/20260706_31MT_S.A._Step5.xlsx'),
    ('S. Pneu','S. pneumoniae','CS1-4_S. Pneu','SPneu','/mnt/user-data/outputs/20260706_31MT_S.Pneu_Step5.xlsx'),
    ('S. Pyo','S. pyogenes','CS1-4_S. Pyo','SPyo','/mnt/user-data/outputs/20260706_31MT_S.Pyo_Step5.xlsx')]

def stability(D):
    cnt=Counter()
    for F in D['FOLDS']:
        tr=F['tr']; ytr=D['Y'][tr]; ntr=int(tr.sum())
        eff,p=ng.eff_filter(F['R'],ytr,ntr); keep=np.where((eff>0.3)&(p<=0.05))[0]
        if len(keep)<1: continue
        rl=ng.group_reps(F['Xtr'][:,keep],F['R'][:,keep],eff[keep],ntr)
        for c in keep[rl]: cnt[D['DESC'][c]]+=1
    return cnt

def comments(short,n1,oq,oa,pQ,pA,bias):
    L=[f'Interpretation ({short}):']
    if n1<=3:
        L+=['!!! WARNUNG: nur %d aktive Objekte. Trotz kleinem p ist dieses Ergebnis NICHT belastbar:'%n1,
            '    bei 3 Aktiven ist die genestete Selektion instabil und der Test durch wenige Extrempunkte',
            '    dominiert. P. aeruginosa bleibt von der praediktiven Modellierung ausgeschlossen (nur deskriptiv).']
    sig = pQ<0.05 and pA<0.05
    if sig and n1>3:
        L+=[f'- Genestet SIGNIFIKANT: Q2={oq:.2f} (p={pQ:.3f}), AUC={oa:.2f} (p={pA:.3f}).',
            '  Auch nach Entfernung des Selektions-Leakage bleibt ein ueberzufaelliges, wenn auch SCHWACHES',
            '  Signal. Das ist ein qualitativer Unterschied zu H. influenzae/S. pyogenes.']
    elif n1>3:
        L+=[f'- Genestet NICHT signifikant: Q2={oq:.2f} (p={pQ:.3f}), AUC={oa:.2f} (p={pA:.3f}).',
            '  Nach Entfernung des Leakage bleibt kein nachweisbares Signal (Q2<=0 bzw. AUC in der Nullwolke).']
    L+=[f'- Verzerrt (Selektion vorab) war Q2={bias[0]:.2f}/AUC={bias[1]:.2f}; der Einbruch quantifiziert das Leakage.',
        '- n:p bleibt kritisch (n=31, %d Aktive); Ergebnis hypothesengenerierend, kein prognostisches Modell.'%n1,
        '- MULTIPLIZITAET: 5 Endpunkte getestet -> bei alpha=0.05 waere ein Treffer zufaellig moeglich;',
        '  Bonferroni (alpha=0.01) bestehen nur Endpunkte mit p<0.01. Bei Berichterstattung angeben.']
    return L

def write(short,full,sheet,code,fpath,D,cnt):
    d=np.load(f'/home/claude/perm_{code}.npz'); q2=d['q2s']; auc=d['aucs']
    oq=float(d['obs_q2']); oa=float(d['obs_auc']); n=len(q2)
    pQ=(1+np.sum(q2>=oq - 1e-9))/(1+n); pA=(1+np.sum(auc>=oa - 1e-9))/(1+n)
    n1=int(D['Y'].sum()); bias=BIASED[short]; stab=sorted(cnt.items(),key=lambda x:-x[1])
    wb=openpyxl.load_workbook(fpath); nm=f'PLS-DA_nested_{short}'
    if nm in wb.sheetnames: del wb[nm]
    ws=wb.create_sheet(nm); H=Font(bold=True,size=12); B=Font(bold=True); fill=PatternFill('solid',fgColor='F0DDDD'); r=1
    def put(v,bold=False,sec=False):
        nonlocal r
        for j,x in enumerate(v,1):
            c=ws.cell(row=r,column=j,value=x)
            if bold:c.font=B
            if sec:c.font=H;c.fill=fill
        r+=1
    put([f'PLS-DA – VOLLSTAENDIG GENESTETE VALIDIERUNG ({full})'],sec=True)
    put(['Effektgroessen-Filter + Feature-Group-Bildung in jeder LOO-Fold nur auf Trainingsdaten;'])
    put([f'{n} Permutationen der Gesamtpipeline. Basis: bereinigter 2208-Deskriptor-Satz.'])
    r+=1
    put(['ERGEBNIS: VERZERRT vs. EHRLICH'],sec=True)
    put(['Kennzahl','verzerrt','ehrlich (genestet)'],bold=True)
    put(['Q2 (LOO)',bias[0],round(oq,3)]); put(['AUC (LOO)',bias[1],round(oa,3)])
    put(['Permutations-p (Q2)','~0.0005',round(pQ,4)]); put(['Permutations-p (AUC)','~0.0005',round(pA,4)])
    r+=1
    put([f'PERMUTATIONSTEST GENESTET ({n} Permutationen)'],sec=True)
    put(['Statistik','Beobachtet','Perm-Median','p-Wert'],bold=True)
    put(['Q2 (LOO)',round(oq,3),round(float(np.median(q2)),3),round(pQ,4)])
    put(['AUC (LOO)',round(oa,3),round(float(np.median(auc)),3),round(pA,4)])
    put(['p = (1 + #{Perm >= Beobachtet}) / (1 + n_perm)'])
    r+=1
    put(['SELEKTIONS-STABILITAET (in wie vielen der 31 Folds gewaehlt)'],sec=True)
    put(['Descriptor','Folds (von 31)'],bold=True)
    for desc,c in stab[:20]: put([desc,c])
    r+=1
    put(['KOMMENTARE'],sec=True)
    for line in comments(short,n1,oq,oa,pQ,pA,bias): ws.cell(row=r,column=1,value=line); r+=1
    ws.column_dimensions['A'].width=46
    for col in ['B','C','D']: ws.column_dimensions[col].width=20
    wb.save(fpath)
    return dict(short=short,n1=n1,oq=oq,oa=oa,pQ=pQ,pA=pA)

res={}
for short,full,sheet,code,fpath in EP:
    D=ng.build(sheet); cnt=stability(D); res[short]=write(short,full,sheet,code,fpath,D,cnt)
    print(f'{short}: geschrieben | n1={res[short]["n1"]} Q2={res[short]["oq"]:.3f} AUC={res[short]["oa"]:.3f} p(Q2)={res[short]["pQ"]:.4f} p(AUC)={res[short]["pA"]:.4f}')
np.save('/home/claude/nested_summary.npy',res)
print('OK')
